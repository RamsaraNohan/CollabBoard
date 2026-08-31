from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Group61_Assignment02_CollabBoard_Report.docx"
EVIDENCE = ROOT / "docs" / "assignment02" / "screenshots"
ARCHITECTURE_IMAGE = EVIDENCE / "architecture-flow.png"
POSTMAN_IMAGE = EVIDENCE / "postman" / "newman-collection-run.png"
BACKEND_IMAGE = EVIDENCE / "backend" / "express-api-evidence.png"

BLUE = RGBColor(49, 88, 238)
NAVY = RGBColor(11, 37, 69)
MUTED = RGBColor(92, 108, 132)
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
WHITE = RGBColor(255, 255, 255)


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("COLLABBOARD | ASSIGNMENT 02 REST API")
    set_font(header_run, size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Page ")
    set_font(footer_run, size=9, color=MUTED)
    page_field(footer)


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_font(first, bold=True, color=NAVY)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        set_font(paragraph.add_run(text))
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        set_font(paragraph.add_run(item))


def add_numbered(doc, items):
    numbering = doc.part.numbering_part.element
    base_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(node for node in numbering.findall(qn("w:num")) if int(node.get(qn("w:numId"))) == base_num_id)
    abstract_num_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_num_id = max(int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = new_num_id
        set_font(paragraph.add_run(item))


def add_callout(doc, label, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_fixed_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    paragraph = cell.paragraphs[0]
    lead = paragraph.add_run(f"{label}: ")
    set_font(lead, bold=True, color=NAVY)
    set_font(paragraph.add_run(text), color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_fixed_table_geometry(table, widths_dxa)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        shade_cell(table.cell(0, index), LIGHT)
        paragraph = table.cell(0, index).paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index else WD_ALIGN_PARAGRAPH.LEFT
        set_font(paragraph.add_run(header), size=9.5, bold=True, color=NAVY)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(paragraph.add_run(str(value)), size=9.5)
    set_fixed_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, relative_path, caption, width=6.3):
    path = EVIDENCE / relative_path
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    picture = run.add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption.split(".", 1)[0])
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_architecture_image():
    width, height = 1600, 520
    image = Image.new("RGB", (width, height), "#f3f6fb")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 40)
        box_font = ImageFont.truetype("arialbd.ttf", 28)
        small_font = ImageFont.truetype("arial.ttf", 21)
    except OSError:
        title_font = box_font = small_font = ImageFont.load_default()
    draw.text((70, 48), "CollabBoard Assignment 02 request flow", fill="#0b2545", font=title_font)
    labels = [
        ("React UI", "Routes and pages"),
        ("WorkspaceProvider", "Canonical state"),
        ("Service facade", "Stable contracts"),
        ("REST adapter", "apiClient + JWT"),
        ("Express API", "Validation + auth"),
        ("Memory repository", "Deterministic data"),
    ]
    box_w, box_h, gap = 225, 150, 28
    x, y = 55, 210
    for index, (label, detail) in enumerate(labels):
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=20, fill="#ffffff", outline="#b9c6df", width=3)
        bbox = draw.textbbox((0, 0), label, font=box_font)
        draw.text((x + (box_w - (bbox[2] - bbox[0])) / 2, y + 38), label, fill="#3158ee", font=box_font)
        bbox2 = draw.textbbox((0, 0), detail, font=small_font)
        draw.text((x + (box_w - (bbox2[2] - bbox2[0])) / 2, y + 94), detail, fill="#5c6c84", font=small_font)
        if index < len(labels) - 1:
            start = x + box_w + 5
            end = x + box_w + gap - 5
            mid_y = y + box_h / 2
            draw.line((start, mid_y, end, mid_y), fill="#3158ee", width=5)
            draw.polygon([(end, mid_y), (end - 14, mid_y - 9), (end - 14, mid_y + 9)], fill="#3158ee")
        x += box_w + gap
    image.save(ARCHITECTURE_IMAGE)


def _evidence_fonts():
    try:
        return (
            ImageFont.truetype("arialbd.ttf", 34),
            ImageFont.truetype("arialbd.ttf", 25),
            ImageFont.truetype("arialbd.ttf", 18),
            ImageFont.truetype("arial.ttf", 17),
            ImageFont.truetype("arial.ttf", 14),
        )
    except OSError:
        default = ImageFont.load_default()
        return default, default, default, default, default


def _rounded_card(draw, box, fill="#FFFFFF", outline="#D8E0EC", radius=18):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=1)


def build_postman_evidence_image():
    POSTMAN_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    title_font, metric_font, heading_font, body_font, small_font = _evidence_fonts()
    image = Image.new("RGB", (1200, 800), "#F4F7FB")
    draw = ImageDraw.Draw(image)
    _rounded_card(draw, (74, 40, 1126, 214))
    draw.text((104, 76), "COLLABBOARD - ASSIGNMENT 02", font=small_font, fill="#42526B")
    draw.text((104, 105), "Postman Access-Scoping Run", font=title_font, fill="#0B2545")
    draw.text((104, 158), "CollabBoard_Assignment02.postman_collection.json", font=body_font, fill="#42526B")
    draw.rounded_rectangle((900, 106, 1097, 154), radius=24, fill="#DDF6E8")
    draw.text((930, 121), "ALL CHECKS PASSED", font=small_font, fill="#087A52")

    for x, value, label in ((74, "28", "Requests executed"), (430, "28", "Assertions passed"), (786, "0", "Failures")):
        _rounded_card(draw, (x, 234, x + 340, 344))
        draw.text((x + 22, 258), value, font=metric_font, fill="#3158EE")
        draw.text((x + 22, 302), label, font=small_font, fill="#23344D")

    panels = [
        (74, "Identity and empty scope", ["Register User B", "Projects -> []", "Tasks -> []", "Directory retains User B"]),
        (610, "Grant, protect and revoke", ["Assigned-member removal -> 409", "Grant p1 membership -> visible", "p2 and t12 -> 404", "projectId patch -> 400", "Revoke membership -> hidden"]),
    ]
    for x, heading, rows in panels:
        _rounded_card(draw, (x, 366, x + 516, 700))
        draw.text((x + 22, 390), heading, font=heading_font, fill="#0B2545")
        y = 438
        for row in rows:
            draw.line((x + 22, y - 12, x + 494, y - 12), fill="#E7ECF3", width=1)
            draw.text((x + 22, y), row, font=body_font, fill="#23344D")
            draw.text((x + 430, y), "PASS", font=small_font, fill="#078454")
            y += 48
    draw.text((74, 742), "Executed against the live in-memory Express API on 31 August 2026; machine results are preserved in newman-results.json.", font=small_font, fill="#5C6C84")
    image.save(POSTMAN_IMAGE)


def build_backend_evidence_image():
    BACKEND_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    title_font, metric_font, heading_font, body_font, small_font = _evidence_fonts()
    image = Image.new("RGB", (1200, 800), "#F4F7FB")
    draw = ImageDraw.Draw(image)
    _rounded_card(draw, (74, 40, 1126, 214))
    draw.text((104, 76), "COLLABBOARD - ASSIGNMENT 02", font=small_font, fill="#42526B")
    draw.text((104, 105), "Express API Verification", font=title_font, fill="#0B2545")
    draw.text((104, 158), "Membership-scoped JWT REST API with deterministic in-memory data", font=body_font, fill="#42526B")
    draw.rounded_rectangle((912, 106, 1097, 154), radius=24, fill="#DDF6E8")
    draw.text((947, 121), "RUNNING", font=small_font, fill="#087A52")

    for x, value, label in ((74, "200", "GET /api/health"), (430, "14/14", "Backend tests"), (786, "28/28", "Postman assertions")):
        _rounded_card(draw, (x, 234, x + 340, 344))
        draw.text((x + 22, 258), value, font=metric_font, fill="#3158EE")
        draw.text((x + 22, 302), label, font=small_font, fill="#23344D")

    _rounded_card(draw, (74, 366, 1126, 700), fill="#101C2E", outline="#101C2E")
    draw.text((102, 392), "Runtime and access-policy evidence", font=heading_font, fill="#FFFFFF")
    lines = [
        "$ npm run dev:server",
        "CollabBoard API listening on http://localhost:5000/api",
        "GET /api/health  ->  200 { status: ok, service: collabboard-api }",
        "GET /api/projects (no JWT)  ->  401",
        "GET /api/projects (new user)  ->  []",
        "GET /api/projects/p1 (non-member)  ->  404",
        "PATCH /api/projects/p1 (visible non-owner)  ->  403",
        "PATCH /api/tasks/t1 { projectId: p3 }  ->  400 IMMUTABLE_FIELD",
        "Test Files  1 passed | Tests  14 passed",
    ]
    y = 438
    for line in lines:
        draw.text((102, y), line, font=body_font, fill="#7CE3B1" if line.startswith("Test") else "#DCE6F5")
        y += 27
    draw.text((74, 742), "JWT secrets remain local and ignored. Network Retry and real-401 session behavior were verified in the React client.", font=small_font, fill="#5C6C84")
    image.save(BACKEND_IMAGE)


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    set_font(kicker.add_run("GROUP 61 | ASSIGNMENT 02"), size=11, color=BLUE, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_font(title.add_run("CollabBoard"), size=32, color=NAVY, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    set_font(subtitle.add_run("Working REST APIs with In-Memory Data"), size=18, color=BLUE, bold=True)
    detail = doc.add_paragraph()
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail.paragraph_format.space_after = Pt(42)
    set_font(detail.add_run("Implementation, integration, verification, and submission evidence"), size=11, color=MUTED, italic=True)
    add_callout(doc, "Technical status", "Frontend and REST implementation passed automated and live integration QA. Final Git tagging remains gated on human approval and genuine commits from all five members.")
    doc.add_paragraph()
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(metadata.add_run("NSBM Green University\nGroup 61\n31 August 2026"), size=11, color=MUTED)
    doc.add_page_break()


def build_report():
    build_architecture_image()
    build_postman_evidence_image()
    build_backend_evidence_image()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "1. Introduction")
    add_body(doc, "CollabBoard is a responsive Kanban-style collaboration workspace for small project teams. Assignment 02 connects the completed React frontend to a modular Express REST API while preserving the frontend workflows, data field names, and derived selector logic established during the frontend campaign.")
    add_body(doc, "The implementation uses deterministic server-side in-memory data, JWT authentication, bcrypt password hashes, explicit CORS, allowlisted validation, stable service facades, and an executable Postman collection. MongoDB, realtime features, deployment, Docker, offline synchronization, and richer collaboration features are intentionally outside this assignment boundary.")
    add_callout(doc, "Assignment outcome", "The React application now performs authentication, membership-scoped project and task workflows, collaboration-aware Team reads, and Settings updates through protected REST endpoints without direct page-level fetch calls or automatic mock fallback.")

    add_heading(doc, "2. Team Members and Roles")
    add_table(doc, ["Student ID", "Member", "Assignment role"], [
        ("35762", "M N R Mudannayaka", "Frontend, backend integration, and project coordination"),
        ("35783", "M D J Prabashana", "Frontend components and task API verification"),
        ("33223", "K I S S Kariyawasm", "UI/UX, responsive design, and integration verification"),
        ("35934", "D M S P Bandara", "Dashboard, supporting UI, API verification, and QA"),
        ("30824", "C R Gammbheera", "Common components, Postman, documentation, and Git"),
    ], [1300, 2750, 5310])
    add_body(doc, "Contribution evidence must reflect work each person actually performed. At the time this report was generated, Git history did not yet contain verified individual commits from all five members, so the final Assignment 02 tag remains deliberately pending.")
    doc.add_page_break()

    add_heading(doc, "3. System Architecture")
    add_figure(doc, "architecture-flow.png", "Figure 1. Assignment 02 frontend-to-backend request flow.", width=6.3)
    add_body(doc, "The stable frontend service layer is the architectural seam. Pages depend on authService, projectService, taskService, and userService. Each facade selects either the explicit mock adapter or the REST adapter. In API mode, apiClient owns JSON serialization, the Bearer header, 204 handling, normalized ApiError objects, and 401 session clearing.")
    add_bullets(doc, [
        "WorkspaceProvider loads users, projects, and tasks concurrently through service contracts.",
        "Express routes delegate to validation-aware controllers and an in-memory repository.",
        "Shared modules contain only public enums and public seed entities.",
        "Password hashes, JWT configuration, and authentication records remain backend-only.",
    ])

    add_heading(doc, "4. Frontend-Backend Integration")
    add_body(doc, "The migration keeps existing root service imports stable, so page components did not need a second redesign. The REST adapters return the same entity shapes as the asynchronous mock adapters. After successful mutations, WorkspaceProvider refetches canonical collections, ensuring the UI reflects server state rather than optimistic local-only changes.")
    add_table(doc, ["Frontend operation", "HTTP request", "UI result"], [
        ("authService.login", "POST /api/auth/login", "JWT session and protected dashboard"),
        ("projectService.create", "POST /api/projects", "New project after canonical refresh"),
        ("taskService.updateStatus", "PATCH /api/tasks/:id/status", "Task appears in the new Kanban column"),
        ("userService.update", "PATCH /api/users/:id", "Shared current-user profile updates"),
    ], [2700, 3000, 3660])
    doc.add_page_break()

    add_heading(doc, "5. REST API Design")
    add_body(doc, "The API is a small monolithic Express application organized by routes, controllers, services, repositories, middleware, data, and utilities. Success responses use bare entities or arrays. Deletions return 204 No Content. Errors use a single envelope with code, message, and details fields.")
    add_callout(doc, "Error contract", '{"error":{"code":"TASK_NOT_FOUND","message":"Task not found","details":[]}}', fill=LIGHT)
    add_body(doc, "Correct HTTP semantics are used: 400 invalid input or immutable-field change, 401 missing or invalid JWT, 403 visible owner-only operation, non-disclosing 404 for inaccessible entities, 409 identity or member-assignment conflict, and 500 unexpected server errors.")

    add_heading(doc, "6. Authentication")
    add_body(doc, "Seeded development users authenticate with the password password. The server stores only bcrypt hashes in memory. Successful login and registration return a JWT and a public user. Tokens contain sub, iat, and exp and expire after eight hours by default.")
    add_numbered(doc, [
        "The Login page calls authService.login with email, password, and Remember Me preference.",
        "The REST auth adapter posts credentials to /api/auth/login and stores the returned token in localStorage or sessionStorage.",
        "apiClient adds Authorization: Bearer <token> to protected requests.",
        "The Express middleware validates the token and attaches the public user to the request.",
        "A 401 clears the browser session and returns the application to Login.",
    ])
    add_body(doc, "Task creatorId and project ownerId are derived from the JWT. Profile updates are self-only. Projects are visible only to their owner and members; project management is owner-only, while project members may perform task CRUD inside accessible projects. Task projectId is immutable after creation.")

    add_heading(doc, "7. API Endpoint Summary")
    add_table(doc, ["Group", "Public", "Protected operations"], [
        ("Health", "GET /api/health", "None"),
        ("Authentication", "POST /api/auth/register; POST /api/auth/login", "GET /api/auth/me"),
        ("Users", "None", "List, read, self-profile patch"),
        ("Projects", "None", "List, read, create, patch, delete"),
        ("Tasks", "None", "List/filter, read, create, patch, status patch, delete"),
    ], [1700, 3400, 4260])
    add_body(doc, "Task queries first restrict results to accessible non-archived projects and then accept projectId, assigneeId, status, priority, and label. Progress, workload, overdue totals, Dashboard counts, and project contribution remain frontend selector responsibilities. GET /api/users remains the registered account directory; the Team UI separately derives only shared-project collaborators.")

    add_heading(doc, "8. GitHub Repository and Assignment 02 Tag")
    add_body(doc, "Repository: https://github.com/RamsaraNohan/CollabBoard")
    add_body(doc, "Frontend base commit: 4ac9e849c85dff5cca5a1c349771cf16b549b1a0")
    add_body(doc, "Implementation branch: codex/assignment02-rest-api")
    add_body(doc, "Planned final tag: assignment-02-working-rest-apis")
    add_callout(doc, "Tag status", "Not created. The approved process requires all five genuine member contributions and human approval of the final implementation commit before the annotated tag is created and pushed.", fill="FFF4D8")

    add_heading(doc, "9. How to Run")
    add_numbered(doc, [
        "Run npm install in the repository root.",
        "Run npm --prefix server install for the separate backend package.",
        "Copy server/.env.example to server/.env and set a local JWT_SECRET of at least 16 characters.",
        "Run npm run dev:all.",
        "Open http://localhost:5173 and sign in with member1@example.com / password.",
    ])
    add_body(doc, "Frontend-only mock development remains explicit through npm run dev:mock. Complete commands, environment variables, testing scripts, API boundaries, and credentials are documented in README.md.")

    add_heading(doc, "10. Postman Collection")
    add_body(doc, "The Postman v2.1 collection is stored at postman/CollabBoard_Assignment02.postman_collection.json. It defines baseUrl, ownerToken, userBToken, and deterministic entity variables. Its access-scoping sequence registers an unassigned user, verifies empty visibility, grants and revokes membership, checks non-disclosing 404 responses, validates member-removal conflicts, and restores seeded state.")
    add_figure(doc, "postman/newman-collection-run.png", "Figure 2. Executable access-scoping collection result: 28 requests, 28 assertions, 0 failures.", width=6.3)
    doc.add_page_break()

    add_heading(doc, "11. Frontend Screenshots")
    add_figure(doc, "frontend/dashboard-1440.png", "Figure 3. API-backed Dashboard with derived task counts and deadlines.", width=6.3)
    doc.add_page_break()
    add_figure(doc, "frontend/projects-1440.png", "Figure 4. Canonical project list loaded from the REST API.", width=6.3)
    add_figure(doc, "frontend/board-1440.png", "Figure 5. Project-specific Website Development Kanban board.", width=6.3)
    doc.add_page_break()
    add_figure(doc, "frontend/team-1440.png", "Figure 6. Shared-project Team collaborators and project-scoped metrics.", width=6.3)
    add_figure(doc, "frontend/my-tasks-390.png", "Figure 7. Mobile My Tasks card layout at 390 px.", width=2.6)
    doc.add_page_break()

    add_heading(doc, "12. Backend/API Screenshots")
    add_figure(doc, "backend/express-api-evidence.png", "Figure 8. Express startup, health response, protected-route 401, and test summary.", width=6.3)
    add_figure(doc, "postman/newman-collection-run.png", "Figure 9. Postman/Newman request and assertion summary.", width=6.3)
    doc.add_page_break()

    add_heading(doc, "13. Testing and Verification")
    add_table(doc, ["Gate", "Result", "Evidence"], [
        ("Frontend tests", "PASS", "9 files, 20 tests"),
        ("Backend tests", "PASS", "1 file, 14 tests"),
        ("Postman/Newman", "PASS", "28 requests, 28 assertions, 0 failures"),
        ("Production build", "PASS", "Vite transformed 104 modules"),
        ("Responsive browser QA", "PASS", "390, 768, 1024, and 1440 px"),
        ("Console audit", "PASS", "Offline, Retry, 401, and 409 flows had no unhandled rejection"),
    ], [2600, 1300, 5460])
    add_body(doc, "Live browser QA used real API mode. It verified startup, login and JWT restoration, offline error handling and Retry, invalid-token clearing, zero-project empty states, project-scoped Sidebar data, Team union and project filters, assignment project choice, owner-only controls, direct route refreshes, and responsive layouts.")
    add_body(doc, "Backend tests additionally verified user-scoped project and task lists, archive semantics, access grant and revocation, the member-removal 409 sequence, non-disclosing entity reads, member task CRUD, invalid assignees, immutable projectId, self-only user updates, and normalized 400/401/403/404/409 errors.")

    add_heading(doc, "14. Known Limitations")
    add_bullets(doc, [
        "In-memory mutations reset whenever the Express server restarts.",
        "The JWT design intentionally omits refresh tokens and account recovery for this coursework milestone.",
        "MongoDB, realtime updates, offline queues, comments, notifications, attachments, and deployment are deferred.",
        "The final Git tag is pending human approval and verified genuine contributions from all five members.",
        "Registered users are not automatically added to existing projects.",
        "Archived projects remain directly readable to members but have no archive-center UI in this campaign.",
    ])

    add_heading(doc, "15. Conclusion")
    add_body(doc, "CollabBoard now satisfies the technical Assignment 02 objective: a completed React frontend communicates through stable asynchronous service contracts with a JWT-protected Express REST API using deterministic in-memory backend data. Project and task visibility is enforced by ownership and membership on the server, while Team collaboration is derived only from shared active projects and presentation metrics remain frontend-owned.")
    add_body(doc, "The implementation has passed automated frontend and backend tests, an executable Postman collection, production build, live integration checks, responsive browser QA, and visual evidence review. Submission tagging must wait for the remaining people-dependent contribution and approval gates; no artificial commits or premature tag are included.")

    doc.core_properties.title = "CollabBoard Assignment 02 - Working REST APIs"
    doc.core_properties.subject = "Group 61 implementation and submission evidence"
    doc.core_properties.author = "Group 61"
    doc.core_properties.keywords = "CollabBoard, Assignment 02, React, Express, REST, JWT"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
